"""Leitura somente-leitura de planilhas (valores em cache, formulas NAO sao
recalculadas) e conversores tolerantes para os formatos sujos encontrados."""
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from fnmatch import fnmatchcase
from pathlib import Path

from openpyxl import load_workbook

MAX_LINHAS_PROCURA_CABECALHO = 30
MARCADORES_VAZIOS = {"", "-", "--", "—", "N/A", "NA", ";", "NULL"}


# ---------------------------------------------------------------------------
# Normalizacao de textos
# ---------------------------------------------------------------------------
def sem_acentos(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c))


def normalizar_cabecalho(valor) -> str | None:
    """`% EXEC`, `%EXEC`, `Qtde. NECs Faturada`, `ÁREA Construção/Manutenção`
    -> comparaveis: sem acento, MAIUSCULO, sem pontuacao, espacos colapsados."""
    if valor is None:
        return None
    t = sem_acentos(str(valor)).upper().replace("%", " % ")
    t = re.sub(r"[^A-Z0-9%]+", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t or None


def normalizar_chave(valor) -> str:
    """Chave de comparacao para nomes (clientes, equipes): sem acento, maiusculo, espacos colapsados."""
    return re.sub(r"\s+", " ", sem_acentos(str(valor)).upper()).strip()


def texto_limpo(valor) -> str | None:
    if valor is None:
        return None
    t = re.sub(r"\s+", " ", str(valor).replace("\xa0", " ")).strip()
    return None if t.upper() in MARCADORES_VAZIOS else t


# ---------------------------------------------------------------------------
# Conversores
# ---------------------------------------------------------------------------
def para_decimal(valor) -> Decimal | None:
    """Numero real, ou texto pt-BR (`R$ 1.484,10`). `#N/A`, `;`, `-`, vazio -> None."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float, Decimal)):
        try:
            d = Decimal(str(valor))
        except InvalidOperation:
            return None
        return d if d.is_finite() else None
    s = str(valor).replace("\xa0", " ").strip()
    if s.upper() in MARCADORES_VAZIOS or s.startswith("#"):
        return None
    s = s.replace("R$", "").replace(" ", "")
    negativo = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        d = Decimal(s)
    except InvalidOperation:
        return None
    return -d if negativo else d


def valor_e_texto(valor) -> bool:
    """True se a celula veio como texto (nao numerica) -- usado na conciliacao dos totais."""
    return isinstance(valor, str)


def para_data(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str):
        s = valor.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
    return None


def para_hora(valor) -> str | None:
    if isinstance(valor, datetime):
        valor = valor.time()
    if isinstance(valor, time):
        return valor.strftime("%H:%M")
    if isinstance(valor, (int, float)) and 0 <= valor < 1:  # fracao do dia
        minutos = round(valor * 24 * 60)
        return f"{minutos // 60:02d}:{minutos % 60:02d}"
    return texto_limpo(valor)


def para_bool_sim_nao(valor) -> bool | None:
    if valor is None:
        return None
    t = sem_acentos(str(valor)).strip().upper()
    if t in ("SIM", "S", "TRUE", "1"):
        return True
    if t in ("NAO", "N", "FALSE", "0"):
        return False
    return None


def jsonavel(valor):
    """Converte uma celula para algo serializavel em JSON; None se vazia."""
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.isoformat(timespec="minutes") if (valor.hour or valor.minute) else valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, time):
        return valor.strftime("%H:%M")
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, (int, float, bool)):
        return valor
    t = texto_limpo(valor)
    return t


# ---------------------------------------------------------------------------
# Planilha / aba / cabecalho
# ---------------------------------------------------------------------------
class AbaNaoEncontrada(Exception):
    pass


class CabecalhoNaoEncontrado(Exception):
    pass


@dataclass
class Tabela:
    """Uma aba com o cabecalho ja localizado. `linhas()` entrega (numero_linha_excel, tupla)."""

    aba: str
    linha_cabecalho: int  # numero Excel (1-based)
    cabecalhos_norm: list[str | None]
    cabecalhos_orig: list[str | None]
    brutas: list[tuple]
    primeira_linha_dados: int  # indice em `brutas`
    mapeadas: set[int] = field(default_factory=set)

    def linhas(self):
        for i in range(self.primeira_linha_dados, len(self.brutas)):
            yield i + 1, self.brutas[i]

    def coluna(self, aliases: list[str]) -> int | None:
        """Primeira coluna cujo cabecalho normalizado casa (exato ou curinga `*`) com algum alias,
        respeitando a ordem de prioridade dos aliases."""
        for alias in aliases:
            for idx, h in enumerate(self.cabecalhos_norm):
                if h and (h == alias or ("*" in alias and fnmatchcase(h, alias))):
                    return idx
        return None

    def mapear(self, spec: dict[str, list[str]]) -> dict[str, int | None]:
        out = {}
        for chave, aliases in spec.items():
            idx = self.coluna(aliases)
            out[chave] = idx
            if idx is not None:
                self.mapeadas.add(idx)
        return out

    def nao_mapeadas(self) -> list[str]:
        return [
            str(self.cabecalhos_orig[i]) for i, h in enumerate(self.cabecalhos_norm)
            if h and i not in self.mapeadas
        ]

    def extras(self, linha: tuple, ignorar: set[int] = frozenset()) -> dict:
        """Colunas NAO mapeadas com valor, por cabecalho original (nada e' descartado)."""
        out: dict = {}
        for i, h in enumerate(self.cabecalhos_orig):
            if i in self.mapeadas or i in ignorar or not h or i >= len(linha):
                continue
            v = jsonavel(linha[i])
            if v is None:
                continue
            chave = str(h).strip()
            n = 2
            while chave in out:
                chave = f"{str(h).strip()}_{n}"
                n += 1
            out[chave] = v
        return out


def pegar(linha: tuple, idx: int | None):
    return linha[idx] if idx is not None and idx < len(linha) else None


class Planilha:
    def __init__(self, caminho: str | Path):
        self.caminho = Path(caminho)
        self.nome = self.caminho.name
        self._wb = load_workbook(self.caminho, read_only=True, data_only=True)
        self._abas = {normalizar_cabecalho(n): n for n in self._wb.sheetnames}

    @property
    def abas(self) -> list[str]:
        return list(self._wb.sheetnames)

    def tem_aba(self, *nomes: str) -> bool:
        return any(normalizar_cabecalho(n) in self._abas for n in nomes)

    def _ws(self, nomes):
        for n in nomes:
            real = self._abas.get(normalizar_cabecalho(n))
            if real:
                return real, self._wb[real]
        raise AbaNaoEncontrada(f"aba {' / '.join(nomes)} nao encontrada (abas: {', '.join(self.abas)})")

    def linhas_brutas(self, *nomes: str) -> tuple[str, list[tuple]]:
        real, ws = self._ws(nomes)
        try:
            ws.reset_dimensions()  # dimensoes gravadas erradas fazem o modo read-only ler so parte
        except Exception:  # noqa: BLE001
            pass
        return real, [tuple(r) for r in ws.iter_rows(min_row=1, values_only=True)]

    def tabela(self, nomes: list[str], obrigatorios: list[list[str]]) -> Tabela:
        """Localiza a linha de cabecalho (primeira, entre as 30 iniciais, que contem TODAS as
        colunas obrigatorias; cada obrigatoria e' uma lista de aliases alternativos)."""
        real, brutas = self.linhas_brutas(*nomes)
        melhor: tuple[int, list] | None = None
        for i, linha in enumerate(brutas[:MAX_LINHAS_PROCURA_CABECALHO]):
            norm = [normalizar_cabecalho(c) for c in linha]
            t = Tabela(real, i + 1, norm, [c for c in linha], brutas, i + 1)
            if all(t.coluna(alts) is not None for alts in obrigatorios):
                return t
            candidatos = sum(1 for h in norm if h)
            if melhor is None or candidatos > melhor[0]:
                melhor = (candidatos, [h for h in norm if h])
        vistos = ", ".join((melhor[1][:40] if melhor else [])) or "(nenhum cabecalho identificavel)"
        faltam = [alts[0] for alts in obrigatorios]
        raise CabecalhoNaoEncontrado(
            f"aba '{real}': nao achei linha de cabecalho com as colunas {faltam}. "
            f"Cabecalhos da linha mais promissora: {vistos}"
        )

    def fechar(self):
        self._wb.close()
