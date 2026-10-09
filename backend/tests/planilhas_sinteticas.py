"""Gera planilhas SINTETICAS que reproduzem os problemas medidos no relatorio de analise
(etapa 1): titulos acima do cabecalho, colunas deslocadas, valores em texto pt-BR, #N/A,
familia na coluna Area, codigos malformados, cliente "-", duplicidades etc."""
from datetime import date, datetime, time

from openpyxl import Workbook

A = "DMP/A.SUL.25.00419"       # obra normal (CGO + SOT + Obras)
B = "DAC/A.ABC.26.00024"       # CGO, UPS ORCADO em texto
D = "DMP/A.NOR.25.01111"       # CGO; vai para ARQUIVO MORTO; data_fim vencida
E = "DMP/A.SUL.26.00500"       # so SOT (cliente real)
F = "DMP/A.SUL.26.00501"       # so SOT (mesmo cliente com outra caixa)
H = "DMP/A.ABC.26.00777"       # so CARTEIRA DE OBRAS
G = "DMP/A.SUL.26.00900"       # so EXTRACAO SIGEO (obra criada so com o codigo)
SEM_OBRA = "DMP/A.SUL.26.02176"  # valido, mas inexistente nas fontes de obra
INVALIDO_1 = "DMP/A.SUL25.06647"
INVALIDO_2 = "DMP/A.SUL.25.007392"


def _linhas(ws, inicio, linhas, col_inicial=1):
    for i, linha in enumerate(linhas):
        for j, v in enumerate(linha):
            if v is not None:
                ws.cell(row=inicio + i, column=col_inicial + j, value=v)


def faturamento(caminho, *, sem_aba_medicao=False):
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("CGO")
    ws["A1"] = "CGO - CARTEIRA GERAL DE OBRAS"
    _linhas(ws, 3, [
        ["PROJETO", "CONTRATO", "CIRCUITO", "UPS ORÇADO", "%EXEC", "DATA TÉRMINO", "OBSERVAÇÕES", "DATA ASSINATURA"],
        [A, 4600004258, "CIRC-1", 1234.567, 0.5, date(2027, 12, 1), "obs da CGO", None],
        [B, 4600004058, "CIRC-2", "100,5", 0.25, None, None, None],
        [INVALIDO_1, 4600004258, None, 10, 0, None, None, None],
        [D, 4600004258, None, 50, 0.1, date(2025, 1, 10), None, None],
        [None, None, None, None, None, None, None, None],
    ])

    ws = wb.create_sheet("CARTEIRA SOT")
    _linhas(ws, 1, [
        ["Projeto", "Data Prog", "Status", "% EXEC", "STATUS SAP", "Origem", "NEC", "Cliente"],
        [A, None, "PROGRAMADO", 0.5, "MEDI", "x@enel.com", 999, "-"],
        [E, None, "EXECUTADO", 0.0, "LIB", None, 10, "A 2 TRANSPORTES LTDA"],
        [F, None, None, None, None, None, 7, "a 2 transportes ltda"],
    ])

    ws = wb.create_sheet("ARQUIVO MORTO")
    ws["A1"] = "ARQUIVADOS"
    _linhas(ws, 2, [[D, "DAC/A.ABC.99.99999"], ["texto solto", "DMP/A.SUL25.00001"]])

    if not sem_aba_medicao:
        ws = wb.create_sheet("MEDIÇÃO")
        ws["A1"] = "TOTAL"
        ws["G1"] = 1003.5
        _linhas(ws, 3, [
            ["Item", "Projeto", "Área Construção/Manutenção", "Ciclo de Medição", "FAMÍLIA", "Qtde. NECs Faturada",
             "Valor (R$) Faturado", "Emitida NF", "Data da Emissão", "NECs Orçado", "NECs Inventariado",
             "Divergências Faturado x Inventariado", "Valor Pago por NEC", "Subtotal de valores glosados"],
            [1, A, "Construção", "C1", "AC", 22.12, 1000.5, "Sim", date(2026, 3, 10), 44.24, 20, 2.12, 45.2, 100],
            [2, D, "Manutenção", "C1", "MA", 10, "R$ 1.484,10", "Sim", None, None, None, None, None, None],
            [3, E, "AC", "C2", None, 5, "R$ 500,00", "Não", None, None, None, None, None, None],
            [4, A, "Construção", "C1", "AC", 1, 1, "Não", None, None, None, None, None, None],
            [5, SEM_OBRA, "Construção", "C1", "AC", 1, 2, "Não", None, None, None, None, None, None],
            [6, F, "Construção", "C1", "AC", 3, "#N/A", "Não", None, None, None, None, None, None],
        ])

    ws = wb.create_sheet("INVENTARIO")
    _linhas(ws, 1, [
        ["PROJETO", "FAMILIA", "VALOR ORÇADO", "VALOR INVENTARIO", "DIFERENÇA DE SALDO", "DATA INVENTÁRIO",
         "ENVIADO PARA FATUR. ENEL", "TEC. RESPONSAVEL", "STATUS", "OBSERVAÇÃO", "CICLO DE MEDIÇÃO"],
        [A, "AC", 1000, 990, -10, date(2026, 2, 1), "SIM", "Fulano", "OBRA ENCERRADA", "ok", "C1"],
        [A, "AC", 500, 500, 0, None, None, None, "EM PROCESSO DE FISCALIZAÇÃO", None, "C2"],
        [INVALIDO_2, "AC", 1, 1, 0, None, None, None, "REPROVA ENEL", None, None],
        [SEM_OBRA, "AC", 1, 1, 0, None, None, None, "REPROVA ENEL", None, None],
    ])

    ws = wb.create_sheet("EXECUCAO")
    _linhas(ws, 1, [
        ["DATA", "PROJETO", "CIRCUITO", "FAMÍLIA", "% EXECUTADO", "NEC ORÇADO", "NEC PROGRAMADA", "NEC EXECUTADA",
         "EMPRESA", "STATUS", "ATRASO", "OBSERVAÇÕES", "COLUNA NOVA"],
        [date(2026, 1, 5), A, "CIRC-1", "AC", 0.5, 10, 5, 4.5, "TEES", "EM EXECUÇÃO", 0, "tudo certo", "extra!"],
        [date(2026, 1, 6), SEM_OBRA, "CIRC-9", "AC", 0.1, 1, 1, 0, "TEES", "PROGRAMADO", None, None, None],
    ])
    wb.save(caminho)


def programacao(caminho, *, sem_coluna_equipe=False, sem_aba_programacao=False):
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("CARTEIRA DE OBRAS")
    _linhas(ws, 1, [
        ["Projeto", "NEC", "Cliente", "% EXEC"],
        [H, 3, "ENEL CLIENTE FINAL SA", 0.3],
        [A, 888, "-", 0.5],
    ])

    ws = wb.create_sheet("EXTRAÇÃO SIGEO")
    _linhas(ws, 1, [
        ["Projeto", "Data Programação", "Status Programação", "Tipo Intervenção", "Número PowerON", "Equipamentos",
         "CHI", "Horário Início", "Horário Fim", "Contratada", "CONCATENAR"],
        [G, date(2026, 5, 4), "PROGRAMADO", "LM", "123", "RL-1; RL-2", 12.5, time(8, 30), time(17, 0), "TEES", "G|2026"],
        [INVALIDO_1, date(2026, 5, 4), None, None, None, None, None, None, None, None, None],
    ])

    ws = wb.create_sheet("GERADOR")
    _linhas(ws, 1, [
        ["ID", "STATUS", "PROJETO", "REGIONAL", "ÁREA", "CONTRATADA/PRÓPRIA", "CIRCUITO", "ENDEREÇO", "CARREGAMENTO"],
        [1, "ATIVO", A, "A.SUL", "Sul", "Contratada", "CIRC-1", "Rua X", 0.8],
        [2, "ATIVO", INVALIDO_2, "A.SUL", "Sul", None, None, None, None],
        [3, "ATIVO", "DMP/A.ABC.26.55555", "A.ABC", "ABC", None, None, None, 0.4],
    ])

    if not sem_aba_programacao:
        ws = wb.create_sheet("PROGRAMAÇÃO")
        ws["C1"] = "PROGRAMAÇÃO TEES 2026"
        cab = ["PROJETO", "DATA", "EQUIPE", "STATUS", "% PROG", "PROG.1", "QTD LM"]
        if sem_coluna_equipe:
            cab[2] = "OUTRA COISA"
        d = date(2025, 3, 3)
        _linhas(ws, 6, [
            cab,
            [A, d, "João Silva", "PROGRAMADO", 0.5, "x", 3],
            [A, date(2025, 3, 4), "JOÃO  SILVA", "EXECUTADO", 1.0, None, None],
            [D, d, "Maria", "STATUS ESTRANHO", None, None, None],
            [A, d, "João Silva", "PROGRAMADO", 0.5, None, None],            # duplicidade da 1a
            [None, d, "Maria", None, None, None, None],                       # sem projeto
            [A, date(2025, 3, 5), "Maria", None, 1.5, None, None],            # % PROG > 1
            [A, None, "Maria", None, None, None, None],                       # sem data
            [SEM_OBRA, d, "Maria", None, None, None, None],                   # sem obra
            [A, date(2025, 3, 6), None, None, None, None, None],              # sem equipe
            ["PROG++", d, "Maria", None, None, None, None],                   # fora do padrao
        ], col_inicial=3)
        # 2 colunas vazias a esquerda (A,B) -- como na planilha real (C...CB)
    wb.save(caminho)
