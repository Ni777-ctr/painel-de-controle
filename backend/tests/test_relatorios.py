from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.models.equipe import Equipe
from app.models.estoque import Almoxarifado, Material, MovimentacaoEstoque
from app.models.financeiro import Fatura, Medicao
from app.models.obra import Obra, Programacao
from app.models.usuario import Usuario
from tests.conftest import criar_perfil, criar_usuario


def _login(client, cred):
    r = client.post("/auth/login", json=cred)
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _dados(db_session):
    db = db_session()
    admin_id = db.query(Usuario.id).filter(Usuario.usuario == "admin").scalar()
    outro = Usuario(nome="Maria Lider", usuario="maria", senha_hash="x", perfil_id="administrador")
    db.add(outro)
    db.flush()
    eq1 = Equipe(nome="Alfa", lider_usuario_id=admin_id)
    eq2 = Equipe(nome="Beta", lider_usuario_id=outro.id)
    db.add_all([eq1, eq2])
    db.flush()
    o1 = Obra(wl="=CMD|'/c calc'!A1", descricao="Obra Alfa", status="em_execucao", equipe_id=eq1.id, data_inicio=date(2026, 8, 10), valor_contrato=1000)
    o2 = Obra(wl="OBRA-BETA", descricao="Obra Beta", status="contratada", equipe_id=eq2.id, data_inicio=date(2026, 9, 5), valor_contrato=2000)
    db.add_all([o1, o2])
    db.flush()
    db.add(Programacao(obra_id=o1.id, equipe_id=eq1.id, data=date(2026, 9, 1)))
    db.add(Programacao(obra_id=o2.id, equipe_id=eq2.id, data=date(2026, 9, 20)))
    m = Medicao(obra_id=o1.id, contrato=1000, percentual=10, necs_medidos=50, responsavel_usuario_id=outro.id, status="Aprovada")
    db.add(m)
    db.flush()
    db.add(Fatura(obra_id=o1.id, medicao_id=m.id, valor=100, vencimento=date(2026, 9, 30)))
    db.add(Fatura(obra_id=o2.id, valor=200, vencimento=date(2026, 11, 30)))
    al = Almoxarifado(nome="Central")
    mat = Material(codigo="P1", nome="Poste", unidade="UN")
    db.add_all([al, mat])
    db.flush()
    db.add(MovimentacaoEstoque(almoxarifado_id=al.id, material_id=mat.id, obra_id=o1.id, usuario_id=admin_id, tipo="saida", quantidade=2))
    db.commit()
    ids = {"o1": o1.id, "o2": o2.id, "maria": outro.id, "admin": admin_id}
    db.close()
    return ids


def _planilha(resp):
    ws = load_workbook(BytesIO(resp.content)).active
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def test_exporta_xlsx_com_cabecalho_de_autoria_e_registro(client, admin_headers, db_session):
    _dados(db_session)
    r = client.get("/relatorios/exportar/programacao", params={"formato": "xlsx"}, headers=admin_headers)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"] and r.headers["content-disposition"].endswith('.xlsx"')
    linhas = _planilha(r)
    textos = [str(c) for l in linhas for c in l if c]
    assert any("Gerado por Admin" in t for t in textos)
    assert any(t == "Obra (WL)" for t in textos)
    assert r.headers["x-relatorio-linhas"] == "2"

    hist = client.get("/relatorios/historico", headers=admin_headers).json()
    assert hist["total"] == 1
    item = hist["itens"][0]
    assert item["tipo"] == "programacao" and item["usuario_nome"] == "Admin" and item["linhas"] == 2 and item["tamanho_bytes"] > 0


def test_filtros_por_periodo_obra_e_responsavel(client, admin_headers, db_session):
    ids = _dados(db_session)

    def linhas_dados(tipo, **params):
        r = client.get(f"/relatorios/exportar/{tipo}", params=params, headers=admin_headers)
        assert r.status_code == 200, r.text
        return int(r.headers["x-relatorio-linhas"])

    assert linhas_dados("programacao") == 2
    assert linhas_dados("programacao", data_inicio="2026-09-10") == 1
    assert linhas_dados("programacao", data_fim="2026-09-10") == 1
    assert linhas_dados("programacao", obra_id=ids["o2"]) == 1
    assert linhas_dados("programacao", responsavel_id=ids["maria"]) == 1  # lider da equipe Beta
    assert linhas_dados("obras", responsavel_id=ids["admin"]) == 1
    assert linhas_dados("obras", data_inicio="2026-09-01") == 1
    # medicao: responsavel proprio da medicao; fatura: responsavel da medicao vinculada
    assert linhas_dados("medicoes", responsavel_id=ids["maria"]) == 1
    assert linhas_dados("faturas", responsavel_id=ids["maria"]) == 1
    assert linhas_dados("faturas", data_inicio="2026-10-01") == 1
    assert linhas_dados("estoque_movimentacoes", obra_id=ids["o1"]) == 1
    assert linhas_dados("estoque_movimentacoes", obra_id=ids["o2"]) == 0

    hist = client.get("/relatorios/historico", headers=admin_headers).json()["itens"]
    filtrado = next(h for h in hist if h["tipo"] == "programacao" and h["obra_id"] == ids["o2"])
    assert filtrado["filtros"]["obra_id"] == ids["o2"]


def test_exporta_pdf(client, admin_headers, db_session):
    _dados(db_session)
    r = client.get("/relatorios/exportar/obras", params={"formato": "pdf"}, headers=admin_headers)
    assert r.status_code == 200, r.text
    assert r.content[:5] == b"%PDF-" and r.headers["content-type"] == "application/pdf"
    # PDF vazio (sem registros) tambem e' valido
    r = client.get("/relatorios/exportar/frota", params={"formato": "pdf"}, headers=admin_headers)
    assert r.status_code == 200 and r.content[:5] == b"%PDF-"


def test_celula_com_formula_e_neutralizada(client, admin_headers, db_session):
    _dados(db_session)
    r = client.get("/relatorios/exportar/obras", headers=admin_headers)
    celulas = [c for l in _planilha(r) for c in l if isinstance(c, str)]
    assert "'=CMD|'/c calc'!A1" in celulas
    assert "=CMD|'/c calc'!A1" not in celulas


def test_validacoes_de_parametros(client, admin_headers, db_session):
    _dados(db_session)
    assert client.get("/relatorios/exportar/inexistente", headers=admin_headers).status_code == 404
    assert client.get("/relatorios/exportar/obras", params={"formato": "docx"}, headers=admin_headers).status_code == 422
    assert client.get("/relatorios/exportar/obras", params={"data_inicio": "2026-10-01", "data_fim": "2026-09-01"}, headers=admin_headers).status_code == 422
    assert client.get("/relatorios/exportar/obras", params={"obra_id": 9999}, headers=admin_headers).status_code == 404
    assert client.get("/relatorios/exportar/obras", params={"responsavel_id": 9999}, headers=admin_headers).status_code == 404


def test_rbac_exportacao_e_historico(client, db_session, admin_headers):
    ids = _dados(db_session)
    # sem relatorios:export
    criar_perfil(db_session, perfil_id="sem-export", permissoes=["obras:read"])
    h = _login(client, criar_usuario(db_session, usuario="se", perfil_id="sem-export"))
    assert client.get("/relatorios/exportar/obras", headers=h).status_code == 403
    # com export mas sem leitura do dominio
    criar_perfil(db_session, perfil_id="export-obras", permissoes=["relatorios:export", "obras:read"])
    h2 = _login(client, criar_usuario(db_session, usuario="eo", perfil_id="export-obras"))
    assert client.get("/relatorios/exportar/obras", headers=h2).status_code == 200
    assert client.get("/relatorios/exportar/faturas", headers=h2).status_code == 403
    assert [t["chave"] for t in client.get("/relatorios/tipos", headers=h2).json()] == ["obras"]

    # historico: quem nao tem auditoria:read so ve os proprios; admin ve todos
    client.get("/relatorios/exportar/obras", headers=admin_headers)
    proprios = client.get("/relatorios/historico", headers=h2).json()
    assert proprios["total"] == 1 and {i["usuario_nome"] for i in proprios["itens"]} == {"eo"}
    assert client.get("/relatorios/historico", headers=admin_headers).json()["total"] == 2


def test_exportacao_gera_auditoria_de_exportacao(client, admin_headers, db_session):
    _dados(db_session)
    client.get("/relatorios/exportar/faturas", params={"formato": "pdf", "data_inicio": "2026-09-01"}, headers=admin_headers)
    logs = client.get("/auditoria", params={"categoria": "exportacao"}, headers=admin_headers).json()
    assert len(logs) == 1
    assert logs[0]["acao"] == "exportacao_faturas"
    assert logs[0]["depois"]["formato"] == "pdf" and logs[0]["depois"]["data_inicio"] == "2026-09-01"
